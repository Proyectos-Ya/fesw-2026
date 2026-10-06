"""Exportar a PDF y Excel contra la aplicación real (HdU 19).

Los renderers son los de verdad (ReportLab y openpyxl); los repositorios y el
correo son dobles en memoria.
"""

from contextlib import asynccontextmanager
from datetime import timedelta
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.application.use_cases.exports.export_jobs import CompleteExportJobUseCase
from app.bootstrap.repositories import (
    get_export_job_repo,
    get_matching_result_repo,
    get_quotation_repo,
    get_tender_repo,
)
from app.bootstrap.services import get_export_background
from app.config import settings
from app.domain.entities.tender import Tender
from app.infrastructure.services.exports.background import AsyncioExportBackground
from app.main import app
from app.shared.datetime_utils import utc_now_naive
from tests.unit.application.export_fakes import (
    InMemoryExportJobRepository,
    InMemoryQuotationRepository,
)
from tests.unit.application.fakes import (
    FakeEmailService,
    InMemoryMatchingResultRepository,
    InMemoryTenderRepository,
)

EMPRESA = {
    "rut": "76.123.456-0",
    "legal_name": "Constructora Andes SpA",
    "description": "Empresa de mantención de áreas verdes y obras menores.",
    "regions": ["Metropolitana"],
    "sectors": ["Construcción"],
    "years_experience": 6,
    "num_employees": 20,
}
CORREO = "rep@andes.cl"


class Entorno:
    def __init__(self) -> None:
        self.tenders = InMemoryTenderRepository()
        self.jobs = InMemoryExportJobRepository()
        self.email = FakeEmailService()

        @asynccontextmanager
        async def abrir():
            yield CompleteExportJobUseCase(
                jobs=self.jobs, email=self.email, base_url="https://app.test"
            )

        self.fondo = AsyncioExportBackground(abrir)
        self.tender = Tender(
            code="1057539-228-COT26",
            name="Mantención de áreas verdes",
            status_id=1,
            status_code="publicada",
            published_at=utc_now_naive() - timedelta(days=1),
            closing_at=utc_now_naive() + timedelta(days=10),
            last_change_at=utc_now_naive(),
            buyer_rut="61.980.170-9",
            buyer_name="Municipalidad de Providencia",
            buyer_unit="Operaciones",
            available_amount_clp=5_000_000,
        )
        self.tenders.tenders[self.tender.id] = self.tender


@pytest.fixture
def entorno():
    e = Entorno()
    app.dependency_overrides[get_tender_repo] = lambda: e.tenders
    app.dependency_overrides[get_matching_result_repo] = (
        InMemoryMatchingResultRepository
    )
    app.dependency_overrides[get_quotation_repo] = InMemoryQuotationRepository
    app.dependency_overrides[get_export_job_repo] = lambda: e.jobs
    app.dependency_overrides[get_export_background] = lambda: e.fondo
    return e


async def _representante(api: AsyncClient) -> dict[str, str]:
    sub = "cccccccc-cccc-4ccc-8ccc-cccccccccccc"
    api.directorio_de_identidad.confirmar(sub)
    token = api.claves.token(sub=sub, email=CORREO, user_metadata={"full_name": "Rep"})
    headers = {"Authorization": f"Bearer {token}"}
    assert (await api.post("/suppliers", json=EMPRESA, headers=headers)).status_code == 201
    return headers


async def test_un_pdf_a_tiempo_se_descarga_en_la_misma_respuesta(api: AsyncClient, entorno):
    # Criterio 3, con el PDF real de ReportLab.
    headers = await _representante(api)

    respuesta = await api.post(
        f"/tenders/{entorno.tender.id}/exports", json={"format": "pdf"}, headers=headers
    )

    assert respuesta.status_code == 200
    assert respuesta.content.startswith(b"%PDF")
    assert "licitacion-1057539-228-COT26.pdf" in respuesta.headers["content-disposition"]
    assert entorno.jobs.jobs == {}


async def test_un_excel_con_secciones_elegidas(api: AsyncClient, entorno):
    # Criterios 4 y 5, con el Excel real de openpyxl.
    headers = await _representante(api)

    respuesta = await api.post(
        f"/tenders/{entorno.tender.id}/exports",
        json={"format": "xlsx", "sections": ["hitos", "montos"]},
        headers=headers,
    )

    assert respuesta.status_code == 200
    assert load_workbook(BytesIO(respuesta.content)).sheetnames == ["Montos", "Hitos"]


async def test_si_tarda_avisa_y_termina_con_correo_y_descarga(
    api: AsyncClient, entorno, monkeypatch
):
    # Criterio 8: con umbral 0, cualquier generación "tarda demasiado".
    monkeypatch.setattr(settings, "export_inline_timeout_seconds", 0.0)
    headers = await _representante(api)

    respuesta = await api.post(
        f"/tenders/{entorno.tender.id}/exports", json={"format": "pdf"}, headers=headers
    )

    assert respuesta.status_code == 202
    job_id = respuesta.json()["job_id"]
    assert "segundo plano" in respuesta.json()["message"]

    await entorno.fondo.wait_idle()

    estado = await api.get(f"/exports/{job_id}", headers=headers)
    assert estado.json()["status"] == "ready"
    [correo] = entorno.email.sent
    assert correo.to == CORREO
    assert f"https://app.test/exportaciones/{job_id}" in correo.text_body
    archivo = await api.get(f"/exports/{job_id}/file", headers=headers)
    assert archivo.status_code == 200
    assert archivo.content.startswith(b"%PDF")


async def test_exportar_exige_sesion(api: AsyncClient, entorno):
    respuesta = await api.post(f"/tenders/{entorno.tender.id}/exports", json={"format": "pdf"})

    assert respuesta.status_code == 401
