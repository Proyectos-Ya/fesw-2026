"""Exportación del borrador a Word (HU-20, B6, CA3).

Se abre el `.docx` generado con la misma librería y se verifica la estructura:
títulos, viñetas, vacíos resaltados y los bloques "Revisar" sombreados.
"""

from io import BytesIO
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

from app.domain.entities.proposal import (
    DraftContent,
    DraftParagraph,
    DraftSection,
    DraftSource,
    ProposalDraft,
    ProposalWarning,
)
from app.infrastructure.services.docx_proposal_exporter import DocxProposalExporter
from tests.unit.application.test_score_tender_on_demand import crear_licitacion


def _seccion(*textos: str) -> DraftSection:
    return DraftSection(
        paragraphs=[DraftParagraph.from_ai_text(t, sources=[]) for t in textos]
    )


def _borrador(tecnico: DraftSection | None = None, advertencias=()) -> ProposalDraft:
    borrador = ProposalDraft(supplier_id=uuid4(), tender_id=uuid4(), status="READY")
    borrador.content = DraftContent(
        offer_name=_seccion("Capacitación PAC en Coyhaique"),
        offer_description=DraftSection(
            paragraphs=[
                DraftParagraph.from_ai_text(
                    "Relatores con [[INSERTAR: años de experiencia]] años de trayectoria.",
                    sources=[
                        DraftSource(id="perfil:region:aysen", label="Región: Aysén")
                    ],
                ),
                DraftParagraph.from_ai_text("Duración de 40 horas.", sources=[]),
            ]
        ),
        required_documents=_seccion("Cotización", "Formulario de transferencias"),
        technical_document=tecnico,
    )
    borrador.warnings = list(advertencias)
    return borrador


def _documento(borrador: ProposalDraft):
    tender = crear_licitacion(borrador.tender_id)
    datos = DocxProposalExporter().to_docx(borrador, tender)
    return Document(BytesIO(datos)), tender


def _sombreado(parrafo) -> bool:
    ppr = parrafo._p.pPr
    return ppr is not None and ppr.find(qn("w:shd")) is not None


def _con_estilo(documento, estilo: str) -> list[str]:
    return [p.text for p in documento.paragraphs if p.style.name == estilo]


def test_titulo_y_secciones_con_jerarquia():
    documento, _ = _documento(_borrador())

    assert _con_estilo(documento, "Heading 1") == ["Capacitación PAC en Coyhaique"]
    assert _con_estilo(documento, "Heading 2") == [
        "Descripción de la oferta",
        "Documentos necesarios",
    ]


def test_los_documentos_necesarios_van_como_vinetas():
    documento, _ = _documento(_borrador())

    assert _con_estilo(documento, "List Bullet") == [
        "Cotización",
        "Formulario de transferencias",
    ]


def test_el_documento_tecnico_solo_si_existe():
    sin, _ = _documento(_borrador())
    con, _ = _documento(_borrador(tecnico=_seccion("Metodología de trabajo.")))

    assert "Documento técnico" not in _con_estilo(sin, "Heading 2")
    assert "Documento técnico" in _con_estilo(con, "Heading 2")
    assert any(p.text == "Metodología de trabajo." for p in con.paragraphs)


def test_los_vacios_van_resaltados_y_con_un_bloque_revisar():
    documento, _ = _documento(_borrador())

    # `documento.paragraphs` arma objetos nuevos en cada acceso: una sola lista.
    parrafos = documento.paragraphs
    indice = next(i for i, p in enumerate(parrafos) if p.text.startswith("Relatores"))
    resaltados = [
        r.text
        for r in parrafos[indice].runs
        if r.font.highlight_color == WD_COLOR_INDEX.YELLOW
    ]
    assert resaltados == ["(Por favor, inserte aquí el valor años de experiencia)"]
    revisar = parrafos[indice + 1]
    assert revisar.text == "Revisar: completar años de experiencia."
    assert _sombreado(revisar)


def test_un_parrafo_sin_vacios_no_lleva_bloque_revisar():
    documento, _ = _documento(_borrador())

    textos = [p.text for p in documento.paragraphs]
    siguiente = textos[textos.index("Duración de 40 horas.") + 1]
    assert not siguiente.startswith("Revisar")


def test_las_advertencias_van_al_inicio_como_bloques_revisar():
    advertencia = ProposalWarning(
        requirement_id="req-1",
        text="Las bases exigen: certificación SEC. La empresa declaró no cumplirlo.",
    )
    documento, _ = _documento(_borrador(advertencias=[advertencia]))

    bloques = [
        p for p in documento.paragraphs if p.text.startswith("Revisar: Las bases")
    ]
    assert len(bloques) == 1 and _sombreado(bloques[0])
    textos = [p.text for p in documento.paragraphs]
    assert textos.index(bloques[0].text) < textos.index("Descripción de la oferta")


def test_no_incluye_las_fuentes_internas():
    """Las fuentes son para quien revisa en Chiripa, no para el comprador."""
    documento, _ = _documento(_borrador())

    assert not any("perfil:region:aysen" in p.text for p in documento.paragraphs)
    assert not any("Región: Aysén" in p.text for p in documento.paragraphs)


def test_indica_la_licitacion_y_que_es_un_borrador():
    documento, tender = _documento(_borrador())

    texto = "\n".join(p.text for p in documento.paragraphs)
    assert tender.code in texto
    assert "Borrador" in texto
