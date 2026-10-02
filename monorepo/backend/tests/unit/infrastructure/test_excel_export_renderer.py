from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.application.use_cases.exports.export_snapshot import ExportSection
from app.infrastructure.services.exports.excel_renderer import OpenpyxlExcelRenderer
from tests.unit.infrastructure.export_samples import con_licitacion, snapshot

TODAS = list(ExportSection)


def _abrir(contenido: bytes):
    return load_workbook(BytesIO(contenido))


def _fila_con(hoja: Worksheet, etiqueta: str) -> tuple:
    for fila in hoja.iter_rows(values_only=True):
        if fila and fila[0] == etiqueta:
            return fila
    raise AssertionError(f"No hay una fila '{etiqueta}' en la hoja {hoja.title}")


def _textos(hoja: Worksheet) -> list:
    return [c for fila in hoja.iter_rows(values_only=True) for c in fila if c is not None]


class TestHojas:
    def test_una_hoja_por_seccion_en_orden(self):
        libro = _abrir(OpenpyxlExcelRenderer().render(snapshot(), TODAS))

        assert libro.sheetnames == [
            "Datos generales",
            "Montos",
            "Datos técnicos",
            "Hitos",
            "Análisis IA",
        ]

    def test_una_seccion_desmarcada_no_genera_su_hoja(self):
        # Criterio 5: p. ej. omitir el análisis de la IA.
        secciones = [s for s in TODAS if s is not ExportSection.ANALISIS_IA]

        libro = _abrir(OpenpyxlExcelRenderer().render(snapshot(), secciones))

        assert "Análisis IA" not in libro.sheetnames

    def test_solo_las_fechas_clave(self):
        # Criterio 5: "incluir solo las fechas clave".
        libro = _abrir(OpenpyxlExcelRenderer().render(snapshot(), [ExportSection.HITOS]))

        assert libro.sheetnames == ["Hitos"]

    def test_el_orden_no_depende_de_como_llegan_las_secciones(self):
        libro = _abrir(
            OpenpyxlExcelRenderer().render(
                snapshot(), [ExportSection.HITOS, ExportSection.DATOS_GENERALES]
            )
        )

        assert libro.sheetnames == ["Datos generales", "Hitos"]

    def test_sin_secciones_no_genera_nada(self):
        with pytest.raises(ValueError):
            OpenpyxlExcelRenderer().render(snapshot(), [])


class TestContenido:
    def test_datos_generales(self):
        hoja = _abrir(OpenpyxlExcelRenderer().render(snapshot(), TODAS))["Datos generales"]

        assert _fila_con(hoja, "Código")[1] == "1057539-228-COT26"
        assert _fila_con(hoja, "Organismo")[1] == "Municipalidad de Ñuñoa"
        assert _fila_con(hoja, "Empresa")[1] == "Constructora Andes"

    def test_los_hitos_son_fechas_reales_en_hora_de_chile(self):
        # Criterio 4: fechas que Excel puede ordenar y filtrar, no texto.
        hoja = _abrir(OpenpyxlExcelRenderer().render(snapshot(), TODAS))["Hitos"]

        cierre = _fila_con(hoja, "Cierre de postulación")
        # 18:00 UTC del 8-oct = 15:00 en Chile.
        assert cierre[1] == datetime(2026, 10, 8, 15, 0)
        celda = next(c for c in hoja["B"] if c.value == datetime(2026, 10, 8, 15, 0))
        assert "yyyy" in celda.number_format

    def test_los_montos_son_numeros_con_formato_de_moneda(self):
        # Criterio 4: montos que se pueden sumar en la planilla.
        hoja = _abrir(OpenpyxlExcelRenderer().render(snapshot(), TODAS))["Montos"]

        disponible = _fila_con(hoja, "Monto disponible")
        assert disponible[1] == 5_000_000
        total = _fila_con(hoja, "Total cotización")
        # 10 x 1.500 + 2,5 x 12.990 = 15.000 + 32.475
        assert Decimal(str(total[-1])) == Decimal("47475")

    def test_la_cotizacion_detalla_cada_material(self):
        hoja = _abrir(OpenpyxlExcelRenderer().render(snapshot(), TODAS))["Montos"]

        fertilizante = _fila_con(hoja, "Fertilizante")
        assert fertilizante[1:] == ("saco", 2.5, 12990, 32475)

    def test_sin_cotizacion_lo_dice(self):
        hoja = _abrir(OpenpyxlExcelRenderer().render(snapshot(quotation=None), TODAS))["Montos"]

        assert "Sin cotización registrada" in _textos(hoja)

    def test_datos_tecnicos_con_cantidades_numericas(self):
        hoja = _abrir(OpenpyxlExcelRenderer().render(snapshot(), TODAS))["Datos técnicos"]

        fila = _fila_con(hoja, "Corte de pasto")
        assert fila == ("Corte de pasto", "70111703", "Plazas y bandejones", 12, "Servicio")

    def test_analisis_con_puntaje_recomendacion_y_justificacion(self):
        hoja = _abrir(OpenpyxlExcelRenderer().render(snapshot(), TODAS))["Análisis IA"]

        assert _fila_con(hoja, "Compatibilidad (%)")[1] == 84
        assert _fila_con(hoja, "Recomendación")[1] == "Postular"
        assert "mantención" in _fila_con(hoja, "Justificación")[1]

    def test_sin_analisis_lo_dice(self):
        hoja = _abrir(
            OpenpyxlExcelRenderer().render(snapshot(analysis=None, score_pct=None), TODAS)
        )["Análisis IA"]

        assert "Sin análisis de compatibilidad" in _textos(hoja)


class TestSeguridad:
    def test_un_texto_que_parece_formula_queda_como_texto(self):
        # Datos de Mercado Público o de las bases no deben ejecutarse al abrir.
        peligroso = '=HYPERLINK("http://malo.example","clic")'
        base = con_licitacion(snapshot(), name=peligroso)

        hoja = _abrir(OpenpyxlExcelRenderer().render(base, TODAS))["Datos generales"]

        celda = next(c for c in hoja["B"] if c.value == peligroso)
        assert celda.data_type == "s"
