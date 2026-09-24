from io import BytesIO


def xlsx_to_text(file_bytes: bytes, file_name: str) -> str:
    """Aplana las hojas de un XLSX a texto tabular para enviarlo a un modelo de lenguaje."""
    try:
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
        output = [f"=== DOCUMENTO EXCEL: {file_name} ==="]
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            output.append(f"\n--- Hoja: {sheet} ---")
            for row in ws.iter_rows(values_only=True):
                row_values = [str(val) if val is not None else "" for val in row]
                if any(row_values):
                    output.append(" | ".join(row_values))
        return "\n".join(output)
    except Exception:
        return f"[Documento Excel adjunto: {file_name} (no se pudo parsear el contenido tabular)]"
