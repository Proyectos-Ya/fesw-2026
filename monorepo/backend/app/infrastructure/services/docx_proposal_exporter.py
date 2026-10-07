"""Exportación del documento técnico a Word con python-docx (HU-20, CA3).

El Word lleva **solo el documento técnico**, que es lo único que se sube como
archivo a la Compra Ágil. El nombre, la descripción y los documentos necesarios
se copian desde la pestaña del borrador al formulario de la plataforma, igual
que las advertencias, que se ven allí.

Estructura: título "Documento técnico", una línea con la oferta y la licitación,
y un H2 por cada sección de la plantilla fija. Los vacíos van resaltados y
seguidos de un bloque "Revisar" sombreado que dice qué completar y, si la
sección la tiene, una sugerencia de qué poner (plan 292, §2.8).

Las fuentes de cada párrafo **no** se exportan: sirven para revisar en Chiripa,
pero el archivo es el que la empresa termina enviando al comprador.
"""

import re
from io import BytesIO

from docx import Document
from docx.document import Document as DocumentoWord
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from app.application.services.proposal_exporter import IProposalExporter
from app.domain.entities.proposal import DraftParagraph, ProposalDraft
from app.domain.entities.tender import Tender

# Lo que `render_placeholders` deja visible en el texto.
_VACIO = re.compile(r"(\(Por favor, inserte aquí el valor [^)]*\))")
_FONDO_REVISAR = "FFF2CC"


def _sombrear(parrafo: Paragraph) -> None:
    """python-docx no expone el sombreado de párrafo: se escribe el w:shd a mano."""
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), _FONDO_REVISAR)
    parrafo._p.get_or_add_pPr().append(sombra)


def _revisar(documento: DocumentoWord, texto: str) -> None:
    parrafo = documento.add_paragraph()
    etiqueta = parrafo.add_run("Revisar: ")
    etiqueta.bold = True
    parrafo.add_run(texto)
    _sombrear(parrafo)


def _texto_con_vacios(parrafo: Paragraph, texto: str) -> None:
    for trozo in _VACIO.split(texto):
        if not trozo:
            continue
        run = parrafo.add_run(trozo)
        if _VACIO.fullmatch(trozo):
            run.bold = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _parrafo(
    documento: DocumentoWord,
    parrafo: DraftParagraph,
    sugerencia: str | None = None,
) -> None:
    """Un párrafo y, si tiene vacíos, su bloque "Revisar".

    La sugerencia de la sección solo acompaña a un vacío: un párrafo completo
    no la muestra, porque el Word termina en manos del comprador.
    """
    salida = documento.add_paragraph()
    _texto_con_vacios(salida, parrafo.text)
    if parrafo.placeholders:
        texto = f"completar {', '.join(parrafo.placeholders)}."
        if sugerencia:
            texto = f"{texto} Sugerencia: {sugerencia}"
        _revisar(documento, texto)


class DocxProposalExporter(IProposalExporter):
    def to_docx(self, draft: ProposalDraft, tender: Tender) -> bytes:
        contenido = draft.content
        if contenido is None or contenido.technical_document is None:
            raise ValueError("El borrador no tiene documento técnico que exportar.")
        documento = Document()

        documento.add_heading("Documento técnico", level=1)
        nombre = " ".join(p.text for p in contenido.offer_name.paragraphs).strip()
        encabezado = documento.add_paragraph(
            f"Oferta: {nombre or tender.name}. Compra Ágil {tender.code} — "
            f"{tender.name}. Borrador generado en Chiripa: revisar antes de enviar."
        )
        for run in encabezado.runs:
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

        for seccion in contenido.technical_document.sections:
            documento.add_heading(seccion.title, level=2)
            # La de la IA es específica de la licitación; la fija, el respaldo.
            sugerencia = seccion.hint or seccion.guidance
            for parrafo in seccion.paragraphs:
                _parrafo(documento, parrafo, sugerencia)

        salida = BytesIO()
        documento.save(salida)
        return salida.getvalue()
