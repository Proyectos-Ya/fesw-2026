"""Excel del detalle de una licitación, con openpyxl (HdU 19, criterios 4 y 5)."""

from collections.abc import Callable, Sequence
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.application.use_cases.exports.export_snapshot import (
    ExportSection,
    ExportSnapshot,
)
from app.infrastructure.services.exports.formatting import (
    NO_INFORMADO,
    format_chile,
    mercado_publico_url,
    to_chile,
)

_FORMATO_CLP = '"$"#,##0'
_FORMATO_FECHA = "dd-mm-yyyy hh:mm"
_ENCABEZADO = Font(bold=True, color="FFFFFF")
_FONDO_ENCABEZADO = PatternFill("solid", fgColor="0F766E")
_ETIQUETA = Font(bold=True)

Celda = str | int | float | Decimal | datetime | None


def _numero(valor: Decimal | float) -> int | float:
    # Entero cuando lo es, para que Excel no muestre "12990,00" donde va "12990".
    return int(valor) if valor == int(valor) else float(valor)


def _poner(hoja: Worksheet, fila: int, columna: int, valor: Celda, formato: str | None = None):
    celda = hoja.cell(row=fila, column=columna, value=valor)
    if isinstance(valor, str):
        # openpyxl toma como fórmula todo texto que empiece con "="; un dato de
        # Mercado Público o de las bases no debe ejecutarse al abrir la planilla.
        celda.data_type = "s"
    if formato:
        celda.number_format = formato
    return celda


def _encabezados(hoja: Worksheet, fila: int, titulos: Sequence[str]) -> None:
    for columna, titulo in enumerate(titulos, start=1):
        celda = _poner(hoja, fila, columna, titulo)
        celda.font = _ENCABEZADO
        celda.fill = _FONDO_ENCABEZADO


def _pares(hoja: Worksheet, filas: Sequence[tuple[str, Celda]], desde: int = 1) -> int:
    fila = desde
    for etiqueta, valor in filas:
        _poner(hoja, fila, 1, etiqueta).font = _ETIQUETA
        _poner(hoja, fila, 2, valor if valor not in (None, "") else NO_INFORMADO)
        fila += 1
    return fila


def _anchos(hoja: Worksheet, anchos: Sequence[int]) -> None:
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho


def _datos_generales(hoja: Worksheet, s: ExportSnapshot) -> None:
    t = s.tender
    _pares(
        hoja,
        [
            ("Código", t.code),
            ("Nombre", t.name),
            ("Estado", t.status_code),
            ("Organismo", t.buyer_name),
            ("Unidad compradora", t.buyer_unit),
            ("Región", t.region),
            ("Comuna", t.commune),
            ("Descripción", t.description),
            ("Empresa", s.supplier_name),
            ("Ficha en Mercado Público", mercado_publico_url(t.code)),
            ("Generado el", format_chile(s.generated_at)),
        ],
    )
    for celda in hoja["B"]:
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    _anchos(hoja, [26, 90])


def _montos(hoja: Worksheet, s: ExportSnapshot) -> None:
    _poner(hoja, 1, 1, "Monto disponible").font = _ETIQUETA
    disponible = s.tender.available_amount_clp
    _poner(hoja, 1, 2, disponible if disponible is not None else NO_INFORMADO, _FORMATO_CLP)

    q = s.quotation
    if q is None:
        _poner(hoja, 3, 1, "Sin cotización registrada")
        _anchos(hoja, [40, 18])
        return

    formato = _FORMATO_CLP if q.currency == "CLP" else "#,##0.00"
    _poner(hoja, 3, 1, f"Cotización ({q.currency})").font = _ETIQUETA
    _encabezados(hoja, 4, ["Material", "Unidad", "Cantidad", "Precio unitario", "Subtotal"])
    fila = 5
    for item in q.items:
        _poner(hoja, fila, 1, item.description)
        _poner(hoja, fila, 2, item.unit)
        _poner(hoja, fila, 3, _numero(item.quantity))
        _poner(hoja, fila, 4, _numero(item.unit_price), formato)
        _poner(hoja, fila, 5, _numero(item.subtotal), formato)
        fila += 1
    _poner(hoja, fila, 1, "Total cotización").font = _ETIQUETA
    total = _poner(hoja, fila, 5, _numero(q.total), formato)
    total.font = _ETIQUETA
    _anchos(hoja, [40, 12, 12, 18, 18])


def _items(hoja: Worksheet, s: ExportSnapshot) -> None:
    _encabezados(hoja, 1, ["Ítem", "Código producto", "Descripción", "Cantidad", "Unidad"])
    if not s.tender.items:
        _poner(hoja, 2, 1, "Sin ítems informados")
    for fila, item in enumerate(s.tender.items, start=2):
        _poner(hoja, fila, 1, item.name)
        _poner(hoja, fila, 2, item.product_code)
        _poner(hoja, fila, 3, item.description)
        _poner(hoja, fila, 4, _numero(Decimal(str(item.quantity))))
        _poner(hoja, fila, 5, item.unit_of_measure)
    _anchos(hoja, [40, 18, 50, 12, 16])


def _hitos(hoja: Worksheet, s: ExportSnapshot) -> None:
    _encabezados(hoja, 1, ["Hito", "Fecha (hora de Chile)"])
    for fila, hito in enumerate(s.key_dates, start=2):
        _poner(hoja, fila, 1, hito.label)
        # Fecha real, no texto: así se puede ordenar y filtrar en la planilla.
        _poner(hoja, fila, 2, to_chile(hito.at), _FORMATO_FECHA)
    _anchos(hoja, [30, 22])


def _analisis(hoja: Worksheet, s: ExportSnapshot) -> None:
    a = s.analysis
    puntaje = round(a.compatibility_score) if a is not None else s.score_pct
    fila = 1
    if puntaje is not None:
        fila = _pares(hoja, [("Compatibilidad (%)", puntaje)])
    if a is None:
        _poner(hoja, fila, 1, "Sin análisis de compatibilidad")
    else:
        _pares(
            hoja,
            [
                ("Recomendación", a.recommendation),
                ("Justificación", a.justification),
                ("Actualizado el", format_chile(a.updated_at)),
            ],
            desde=fila,
        )
    for celda in hoja["B"]:
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    _anchos(hoja, [22, 100])


_HOJAS: dict[ExportSection, tuple[str, Callable[[Worksheet, ExportSnapshot], None]]] = {
    ExportSection.DATOS_GENERALES: ("Datos generales", _datos_generales),
    ExportSection.MONTOS: ("Montos", _montos),
    ExportSection.ITEMS: ("Datos técnicos", _items),
    ExportSection.HITOS: ("Hitos", _hitos),
    ExportSection.ANALISIS_IA: ("Análisis IA", _analisis),
}


class OpenpyxlExcelRenderer:
    def render(self, snapshot: ExportSnapshot, sections: Sequence[ExportSection]) -> bytes:
        elegidas = set(sections)
        if not elegidas:
            raise ValueError("Hay que elegir al menos una sección para el Excel.")

        libro = Workbook()
        libro.remove(libro.active)
        # Siempre en el mismo orden, venga como venga la selección.
        for seccion in ExportSection:
            if seccion in elegidas:
                titulo, llenar = _HOJAS[seccion]
                llenar(libro.create_sheet(titulo), snapshot)

        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue()
